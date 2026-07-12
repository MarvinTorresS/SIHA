# tests/test_excel_loader.py
# Suite de pruebas para excel_loader.py — reemplazo de malla_scraper.py
# como fuente de la malla curricular (PLANES_DE_ESTUDIO.xlsx).
#
# Cubre las 6 reglas de transformación acordadas:
#   1. Construcción de ciclo_id "N M" (y el caso especial Nivel="Opt")
#   2. Parseo de Requisitos normales (lista de códigos)
#   3. Detección del patrón "Completar ciclos X a Y" (prerequisito_ciclo)
#   4. Exclusión de generales vía placeholders GEN1-4
#   5. Namespacing de optativas por carrera (OPT{n}_{SIGLA})
#   6. Flag es_optativo explícito (no inferido del código)
#
# Ejecutar: pytest tests/test_excel_loader.py -v

import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).parent.parent))

from extractors.excel_loader import (
    CursoMallaExcel,
    SHEET_TO_SIGLA,
    _cargar_hoja,
    _clasificar_placeholder,
    _es_codigo_placeholder,
    _parse_ciclo_id,
    _parse_creditos,
    _procesar_correquisitos,
    _procesar_requisitos,
    cargar_todas_las_carreras,
)

RUTA_EXCEL_REAL = Path(__file__).parent.parent / "input" / "PLANES_DE_ESTUDIO.xlsx"

HEADER = ("Nivel", "Ciclo", "Código", "Curso", "Créditos", "Requisitos", "Correquisitos", "Verificación")


def _construir_hoja(filas: list[tuple]):
    """Construye una hoja openpyxl en memoria: título, fuente, blanco, header, datos, blanco, notas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(("Carrera de prueba",))
    ws.append(("Fuente de prueba",))
    ws.append((None,))
    ws.append(HEADER)
    for f in filas:
        ws.append(f)
    ws.append((None,))
    ws.append(("Nota al pie que no debe procesarse",))
    return ws


# ══════════════════════════════════════════════════════════════════════════════
# Regla 1: ciclo_id "N M"
# ══════════════════════════════════════════════════════════════════════════════

class TestParseCicloId:

    def test_nivel_i_ciclo_i(self):
        assert _parse_ciclo_id("I", "I") == ("1 1", "Primer Año - I Ciclo")

    def test_nivel_iv_ciclo_ii(self):
        assert _parse_ciclo_id("IV", "II") == ("4 2", "Cuarto Año - II Ciclo")

    def test_nivel_iii_ciclo_i(self):
        assert _parse_ciclo_id("III", "I") == ("3 1", "Tercer Año - I Ciclo")

    def test_tolera_espacios_y_minusculas(self):
        assert _parse_ciclo_id(" ii ", " i ") == ("2 1", "Segundo Año - I Ciclo")

    def test_nivel_invalido_lanza_keyerror(self):
        with pytest.raises(KeyError):
            _parse_ciclo_id("V", "I")

    def test_catalogo_electivas_nivel_opt_usa_ciclo_literal(self):
        # Regla confirmada: Nivel="Opt" (catálogo con código propio, ej. EIF100O)
        # no encaja en el patrón "N M" porque no trae Ciclo — usa el mismo
        # literal que ya usa populate_db.py para generales de Banner.
        filas = [
            ("Opt", "—", "EIF100O", "Dispositivos para comunicaciones de datos", 3, "EIF208", "—", "ok"),
        ]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "Ingeniería en Sistemas Informáticos")
        assert cursos[0].ciclo_id == "Opt"
        assert cursos[0].ciclo_nombre == "Optativos"


# ══════════════════════════════════════════════════════════════════════════════
# Regla 2: Requisitos normales (lista de códigos, AND simple)
# ══════════════════════════════════════════════════════════════════════════════

class TestProcesarRequisitosNormales:

    def test_un_solo_codigo(self):
        prereqs, rango = _procesar_requisitos("MAT030", {})
        assert prereqs == ["MAT030"]
        assert rango is None

    def test_varios_codigos_separados_por_coma(self):
        # Caso real: ISI EIF201 "Programación I"
        prereqs, rango = _procesar_requisitos("EIF200, MAT030", {})
        assert prereqs == ["EIF200", "MAT030"]
        assert rango is None

    def test_guion_sin_requisitos(self):
        prereqs, rango = _procesar_requisitos("—", {})
        assert prereqs == []
        assert rango is None

    def test_celda_none_sin_requisitos(self):
        prereqs, rango = _procesar_requisitos(None, {})
        assert prereqs == []
        assert rango is None

    def test_admision_sin_requisitos(self):
        # Caso real: hoja "Enseñanza del Inglés"
        prereqs, rango = _procesar_requisitos("Admisión", {})
        assert prereqs == []
        assert rango is None

    def test_admision_sin_acento_tambien_se_reconoce(self):
        prereqs, _ = _procesar_requisitos("admision", {})
        assert prereqs == []

    def test_codigos_se_normalizan_a_mayusculas(self):
        prereqs, _ = _procesar_requisitos("eif200, mat030", {})
        assert prereqs == ["EIF200", "MAT030"]


# ══════════════════════════════════════════════════════════════════════════════
# Regla 3: patrón "Completar ciclos X a Y" → prerequisito_ciclo
# ══════════════════════════════════════════════════════════════════════════════

class TestDeteccionRangoCiclos:

    def test_patron_real_topografia_tgf425(self):
        # Caso real exacto: TGF425 "Práctica Profesional Supervisada"
        prereqs, rango = _procesar_requisitos(
            "Completar ciclos 1 1 a 3 2 (no es prerreq. de curso)", {}
        )
        assert prereqs == []
        assert rango == {"desde": "1 1", "hasta": "3 2"}

    def test_patron_sin_texto_final(self):
        prereqs, rango = _procesar_requisitos("Completar ciclos 2 1 a 4 1", {})
        assert prereqs == []
        assert rango == {"desde": "2 1", "hasta": "4 1"}

    def test_no_confunde_lista_de_codigos_con_el_patron(self):
        prereqs, rango = _procesar_requisitos("TGF412, TGF413", {})
        assert rango is None
        assert prereqs == ["TGF412", "TGF413"]

    def test_patron_inesperado_en_correquisitos_no_falla_y_loguea(self, caplog):
        # Regla 3: no debería aparecer acá; si aparece, advertencia + texto crudo,
        # nunca una excepción ni pérdida silenciosa del dato.
        texto = "Completar ciclos 1 1 a 3 2 (no es prerreq. de curso)"
        with caplog.at_level("WARNING"):
            correq = _procesar_correquisitos(texto, "topo/TGF999", {})
        assert correq == [texto]
        assert any("inesperado en Correquisitos" in r.message for r in caplog.records)

    def test_integracion_topografia_tgf425_y_tgf426(self):
        # Extraído tal cual del Excel real (Topografía y Catastro).
        filas = [
            ("IV", "II", "TGF425", "Práctica Profesional Supervisada", 4,
             "Completar ciclos 1 1 a 3 2 (no es prerreq. de curso)", "—", "ok"),
            ("IV", "II", "TGF426", "Métodos y Técnicas de Investigación", 4,
             "Completar ciclos 1 1 a 3 2 (no es prerreq. de curso)", "—", "ok"),
        ]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "topo", "Topografía")
        for c in cursos:
            assert c.prerequisitos == []
            assert c.prerequisito_ciclo == {"desde": "1 1", "hasta": "3 2"}


# ══════════════════════════════════════════════════════════════════════════════
# Regla 4: exclusión de generales (placeholders GEN1-4)
# ══════════════════════════════════════════════════════════════════════════════

class TestExclusionGenerales:

    def test_clasifica_estudios_generales_como_gen(self):
        assert _clasificar_placeholder("—", "Estudios Generales I") == "GEN"

    def test_clasifica_sin_numero_romano_tambien_como_gen(self):
        # Caso real: Administración usa "Estudios Generales" a secas, sin I/II/III/IV.
        assert _clasificar_placeholder("—", "Estudios Generales") == "GEN"

    def test_codigo_real_no_es_placeholder(self):
        assert _clasificar_placeholder("EIF200", "Fundamentos de informática") is None

    def test_genera_gen1_a_gen4_en_orden_de_aparicion(self):
        # Reproduce el patrón real de ISI: 2 generales en "1 1" y 2 en "1 2".
        filas = [
            ("I", "I", "EIF200", "Fundamentos de informática", 3, "—", "—", "ok"),
            ("I", "I", "—", "Estudios Generales I", 3, "—", "—", "ok"),
            ("I", "I", "—", "Estudios Generales II", 3, "—", "—", "ok"),
            ("I", "II", "—", "Estudios Generales III", 3, "—", "—", "ok"),
            ("I", "II", "—", "Estudios Generales IV", 3, "—", "—", "ok"),
        ]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "Ingeniería en Sistemas Informáticos")
        generales = [c for c in cursos if c.codigo.startswith("GEN")]
        assert [c.codigo for c in generales] == ["GEN1", "GEN2", "GEN3", "GEN4"]

    def test_generales_quedan_con_carrera_sigla_correcta(self):
        # populate_db.py los excluye por código igual, pero deben traer
        # carrera_sigla consistente por si se necesita en otro punto del pipeline.
        filas = [("I", "I", "—", "Estudios Generales I", 3, "—", "—", "ok")]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "topo", "Topografía")
        assert cursos[0].codigo == "GEN1"
        assert cursos[0].carrera_sigla == "topo"

    def test_generales_no_quedan_marcados_como_optativos(self):
        filas = [("I", "I", "—", "Estudios Generales I", 3, "—", "—", "ok")]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        assert cursos[0].es_optativo is False

    def test_cantidad_distinta_de_4_solo_loguea_advertencia_no_falla(self, caplog):
        filas = [
            ("I", "I", "—", "Estudios Generales I", 3, "—", "—", "ok"),
            ("I", "I", "—", "Estudios Generales II", 3, "—", "—", "ok"),
        ]
        ws = _construir_hoja(filas)
        with caplog.at_level("WARNING"):
            cursos = _cargar_hoja(ws, "isi", "ISI")
        assert len(cursos) == 2
        assert any("se esperaban 4" in r.message for r in caplog.records)


# ══════════════════════════════════════════════════════════════════════════════
# Regla 5: namespacing de optativas por carrera (OPT{n}_{SIGLA})
# ══════════════════════════════════════════════════════════════════════════════

class TestNamespacingOptativas:

    def test_clasifica_optativa_con_guion_como_opt(self):
        # Caso real: ISI usa "—" (no el literal "OPT") para el slot de optativa.
        assert _clasificar_placeholder("—", "Optativa") == "OPT"

    def test_clasifica_optativa_con_literal_opt_como_opt(self):
        # Caso real: Inglés/Topografía/Ciencia de Datos usan el literal "OPT".
        assert _clasificar_placeholder("OPT", "Optativo libre") == "OPT"

    def test_deteccion_normalizada_mayusculas_y_espacios(self):
        # Pedido explícito del usuario: "OPTATIVA" en mayúsculas o con espacios
        # de más no debe romper la detección silenciosamente.
        assert _clasificar_placeholder("—", "  OPTATIVA   ") == "OPT"
        assert _clasificar_placeholder("opt", "Curso Optativo II (Educología)") == "OPT"

    def test_genera_codigos_namespaced_en_orden(self):
        filas = [
            ("IV", "I", "—", "Optativa", 3, "—", "—", "ok"),
            ("IV", "I", "—", "Optativa", 3, "—", "—", "ok"),
            ("IV", "II", "—", "Optativa", 3, "—", "—", "ok"),
            ("IV", "II", "—", "Optativa", 3, "—", "—", "ok"),
        ]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        assert [c.codigo for c in cursos] == ["OPT1_ISI", "OPT2_ISI", "OPT3_ISI", "OPT4_ISI"]

    def test_mismo_codigo_base_no_colisiona_entre_carreras(self):
        # El bug confirmado en populate_db.py: dedup global por código, ignora
        # carrera. El namespacing debe producir códigos distintos por carrera
        # para que ninguna optativa se pierda en ese dedup.
        fila_isi = [("IV", "I", "—", "Optativa", 3, "—", "—", "ok")]
        fila_topo = [("III", "I", "OPT", "Optativo I Libre", 3, "—", "—", "ok")]
        fila_itecad = [("IV", "I", "OPT", "Optativo Disciplinar #1", 3, "—", "—", "ok")]

        codigos = set()
        for sigla, filas in (("isi", fila_isi), ("topo", fila_topo), ("itecad", fila_itecad)):
            ws = _construir_hoja(filas)
            cursos = _cargar_hoja(ws, sigla, sigla)
            codigos.update(c.codigo for c in cursos)

        assert codigos == {"OPT1_ISI", "OPT1_TOPO", "OPT1_ITECAD"}
        assert len(codigos) == 3  # sin colisión

    def test_optativas_insertables_a_diferencia_de_generales(self):
        # Regla 5 vs 4: las optativas SÍ deben quedar en la lista de cursos
        # (populate_db.py las inserta), a diferencia de GEN1-4.
        filas = [
            ("I", "I", "—", "Estudios Generales I", 3, "—", "—", "ok"),
            ("IV", "I", "—", "Optativa", 3, "—", "—", "ok"),
        ]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        codigos = [c.codigo for c in cursos]
        assert "GEN1" in codigos
        assert "OPT1_ISI" in codigos


# ══════════════════════════════════════════════════════════════════════════════
# Regla 6: es_optativo explícito (no inferido del código)
# ══════════════════════════════════════════════════════════════════════════════

class TestEsOptativoFlag:

    def test_placeholder_opt_generado_queda_marcado(self):
        filas = [("IV", "I", "—", "Optativa", 3, "—", "—", "ok")]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        assert cursos[0].es_optativo is True

    def test_curso_normal_no_queda_marcado(self):
        filas = [("I", "I", "EIF200", "Fundamentos de informática", 3, "—", "—", "ok")]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        assert cursos[0].es_optativo is False

    def test_codigo_catalogo_terminado_en_o_no_se_infiere_como_optativo(self):
        # A diferencia de malla_scraper.CursoMalla.__post_init__ (que marcaba
        # es_optativo=True por endswith("O")), la regla nueva NO infiere del
        # código: EIF100O es una electiva con código propio real (Nivel="Opt"),
        # no un placeholder generado, así que debe quedar es_optativo=False.
        filas = [("Opt", "—", "EIF100O", "Dispositivos para comunicaciones de datos",
                   3, "EIF208", "—", "ok")]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        assert cursos[0].codigo == "EIF100O"
        assert cursos[0].es_optativo is False


# ══════════════════════════════════════════════════════════════════════════════
# Caso confirmado aparte: resolución de prerrequisito por nombre (Administración)
# ══════════════════════════════════════════════════════════════════════════════

class TestResolucionPrerreqPorNombre:

    def test_optativa_iii_resuelve_contra_codigo_sintetico_de_optativa_i(self):
        # Extraído tal cual del Excel real (hoja Administración).
        filas = [
            ("IV", "I", "—", "Optativa I", 3, "—", "—", "ok"),
            ("IV", "I", "—", "Optativa II", 3, "—", "—", "ok"),
            ("IV", "II", "—", "Optativa III", 3, "Optativa I", "—", None),
            ("IV", "II", "—", "Optativa IV", 3, "Optativa II", "—", None),
        ]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "adm", "Administración")
        por_codigo = {c.codigo: c for c in cursos}

        assert por_codigo["OPT1_ADM"].nombre == "Optativa I"
        assert por_codigo["OPT3_ADM"].prerequisitos == ["OPT1_ADM"]
        assert por_codigo["OPT4_ADM"].prerequisitos == ["OPT2_ADM"]

    def test_nombre_sin_match_se_trata_como_codigo_literal(self):
        filas = [("I", "I", "EIF201", "Programación I", 4, "EIF200, MAT030", "—", "ok")]
        ws = _construir_hoja(filas)
        cursos = _cargar_hoja(ws, "isi", "ISI")
        assert cursos[0].prerequisitos == ["EIF200", "MAT030"]


# ══════════════════════════════════════════════════════════════════════════════
# Créditos: valores no numéricos ("—", vacío) no deben romper el pipeline
# ══════════════════════════════════════════════════════════════════════════════

class TestParseCreditos:

    def test_entero(self):
        assert _parse_creditos(3, "ctx") == 3.0

    def test_guion_es_cero(self):
        # Caso real: catálogo de electivas de Topografía (TGF440O y otros)
        # trae créditos "—" porque el brochure no los detalla.
        assert _parse_creditos("—", "ctx") == 0.0

    def test_none_es_cero(self):
        assert _parse_creditos(None, "ctx") == 0.0

    def test_texto_no_numerico_loguea_advertencia_y_usa_cero(self, caplog):
        with caplog.at_level("WARNING"):
            resultado = _parse_creditos("tres", "isi/EIF999")
        assert resultado == 0.0
        assert any("no numéricos" in r.message for r in caplog.records)


# ══════════════════════════════════════════════════════════════════════════════
# Integración contra el archivo real (si está presente en input/)
# ══════════════════════════════════════════════════════════════════════════════

@pytest.mark.skipif(not RUTA_EXCEL_REAL.exists(), reason="PLANES_DE_ESTUDIO.xlsx no está en data-extractor/input/")
class TestCargaArchivoReal:

    @pytest.fixture
    def cursos(self):
        return cargar_todas_las_carreras(RUTA_EXCEL_REAL)

    def test_carga_las_5_carreras(self, cursos):
        siglas = set(c.carrera_sigla for c in cursos)
        assert siglas == set(SHEET_TO_SIGLA.values())

    def test_cada_carrera_tiene_exactamente_4_generales(self, cursos):
        for sigla in SHEET_TO_SIGLA.values():
            generales = [c for c in cursos if c.carrera_sigla == sigla and c.codigo.startswith("GEN")]
            assert len(generales) == 4, f"{sigla}: {len(generales)} generales, se esperaban 4"

    def test_codigos_opt_generados_no_colisionan_entre_carreras(self, cursos):
        # Verifica en los datos reales que el namespacing evita el bug de
        # colisión confirmado en populate_db.py (dedup global por código).
        # Nota: NO se puede pedir unicidad global de *todos* los códigos —
        # cursos de servicio compartidos (LIX410, MAT001, MAT002...) aparecen
        # legítimamente repetidos en varias carreras (es el mismo curso real),
        # y GEN1-4 también se repite a propósito por carrera (populate_db.py
        # los excluye de todos modos). Lo que el namespacing debe garantizar
        # es que los códigos OPT{n}_{SIGLA} generados sí sean únicos.
        opt_codigos = [c.codigo for c in cursos if c.codigo.startswith("OPT") and "_" in c.codigo]
        assert opt_codigos, "no se generó ninguna optativa namespaced"
        assert len(opt_codigos) == len(set(opt_codigos)), "hay códigos OPT*_SIGLA duplicados"

    def test_generales_se_repiten_por_carrera_a_proposito(self, cursos):
        # GEN1-4 se repite igual en cada carrera por diseño (regla 4): no es
        # namespaced porque populate_db.py los excluye de curso de todos modos.
        gen1_por_carrera = {c.carrera_sigla for c in cursos if c.codigo == "GEN1"}
        assert gen1_por_carrera == set(SHEET_TO_SIGLA.values())

    def test_tgf425_y_tgf426_traen_prerequisito_ciclo(self, cursos):
        por_codigo = {c.codigo: c for c in cursos}
        assert por_codigo["TGF425"].prerequisito_ciclo == {"desde": "1 1", "hasta": "3 2"}
        assert por_codigo["TGF426"].prerequisito_ciclo == {"desde": "1 1", "hasta": "3 2"}

    def test_administracion_optativa_iii_resuelve_a_optativa_i(self, cursos):
        por_codigo = {c.codigo: c for c in cursos}
        assert por_codigo["OPT3_ADM"].prerequisitos == ["OPT1_ADM"]

    def test_ningun_curso_optativo_generado_pierde_es_optativo(self, cursos):
        opt_generados = [c for c in cursos if c.codigo.split("_")[0].startswith("OPT") and "_" in c.codigo]
        assert opt_generados, "no se generó ninguna optativa namespaced"
        assert all(c.es_optativo for c in opt_generados)
