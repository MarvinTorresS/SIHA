# excel_loader.py
"""
Loader de malla curricular desde PLANES_DE_ESTUDIO.xlsx
=========================================================
Reemplaza a extractors/malla_scraper.py como fuente de la malla curricular.
La malla ya no viene de scraping a botsrb.info: viene de un Excel verificado
manualmente contra brochures/PDFs oficiales (input/PLANES_DE_ESTUDIO.xlsx),
con una hoja LEEME + una hoja por carrera (columnas: Nivel, Ciclo, Código,
Curso, Créditos, Requisitos, Correquisitos, Verificación — esta última es
solo para uso humano y no se procesa).

Produce output/normalized/syllabi_normalized.json con el mismo esquema que
ya consumía populate_db.py (populate_db.py no se modifica en este cambio):
codigo, nombre, creditos, ciclo_id, ciclo_nombre, carrera_sigla,
carrera_nombre, prerequisitos, es_optativo — más dos campos nuevos que
populate_db.py todavía no lee: correquisitos y prerequisito_ciclo (este
último alimenta la tabla prerequisito_ciclo de la migración V5).

Reglas de transformación (decididas junto al usuario, ver conversación):

  1. ciclo_id "N M": Nivel romano (I-IV) → arábigo, Ciclo (I/II) → (1/2).
     Las filas de catálogo de electivas (Nivel="Opt", sin Ciclo) usan el
     literal ciclo_id="Opt" / ciclo_nombre="Optativos" — mismo valor que
     ya usa populate_db.py para los generales de Banner (líneas 404-405)
     y que api/src/services/malla.service.ts ya ordena al final
     (compareCycleIds).

  2. Requisitos: lista de códigos separada por coma → prerequisitos (AND
     simple). Patrón "Completar ciclos X a Y" → NO es prerrequisito de
     curso, va a prerequisito_ciclo={"desde","hasta"}. "Admisión" o "—"
     → sin prerrequisitos.

  3. Correquisitos: mismo parseo que Requisitos, a un campo separado. Si
     aparece el patrón de rango de ciclos aquí (no debería), se loguea
     como advertencia y se conserva el texto crudo, sin fallar.

  4. Código "—" + nombre que empieza con "estudios generales" (comparado
     normalizado: minúsculas, espacios colapsados): son cursos sin código
     propio en la malla. Se les genera GEN1..GEN4 en el orden en que
     aparecen dentro de la carrera — códigos que populate_db.py ya excluye
     explícitamente vía PLACEHOLDERS_GENERALES (populate_db.py líneas
     174-198), sin necesidad de tocar ese script.

  5. Código "—" u "OPT" (varía por hoja) + nombre que contiene "optativ"
     (Optativa/Optativo/Curso Optativo, comparado normalizado): SÍ deben
     insertarse en curso, a diferencia de los generales. Se genera un
     código namespaced por carrera OPT{n}_{SIGLA} (ej. OPT1_ISI, OPT2_ISI,
     OPT1_ITECAD, OPT1_TOPO) para evitar el bug ya confirmado en
     populate_db.py: el dedup de _insert_courses_from_syllabus (líneas
     ~163-171) deduplica solo por código, ignorando la carrera, así que
     un código "OPT1" compartido entre carreras perdería silenciosamente
     las de todas menos la primera carrera procesada.

     Caso especial (Administración): "Optativa III"/"Optativa IV" traen
     como Requisito el *nombre* de una fila optativa hermana ("Optativa I"
     / "Optativa II"), no un código. Se resuelve contra el código
     sintético ya generado para esa fila dentro de la misma carrera.

  6. es_optativo=True solo para los códigos generados por la regla 5 (no
     se infiere de prefijo/sufijo del código como hacía
     CursoMalla.__post_init__ en malla_scraper.py). Los cursos de catálogo
     de electivas con código propio real (ej. EIF100O, TGF440O, TDC100O —
     filas con Nivel="Opt") NO se marcan es_optativo aunque su código
     termine en "O": bajo esta regla nueva, solo la familia OPT* generada
     cuenta.

Ejecución (desde data-extractor/):
  python -m extractors.excel_loader
"""

import json
import logging
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent.parent))

from config import BASE_DIR, CARRERAS_CAMPUS_COTO, LOGS_DIR, OUTPUT_NORMALIZED_DIR

# ── Logging ───────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler(LOGS_DIR / "excel_loader.log", encoding="utf-8"),
    ],
)
log = logging.getLogger(__name__)

# ── Constantes ────────────────────────────────────────────────────────────────

# Anclado a config.BASE_DIR (data-extractor/), no a Path(__file__), para que
# no dependa de en qué subcarpeta viva este script.
INPUT_DIR = BASE_DIR / "input"
EXCEL_FILENAME = "PLANES_DE_ESTUDIO.xlsx"

# Nombre exacto de la hoja en el Excel → sigla interna (config.CARRERAS_CAMPUS_COTO).
SHEET_TO_SIGLA: dict[str, str] = {
    "ISI": "isi",
    "Enseñanza del Inglés": "ingles",
    "Topografía y Catastro": "topo",
    "Ciencia de Datos": "itecad",
    "Administración": "adm",
}

NIVEL_ROMANO = {"I": 1, "II": 2, "III": 3, "IV": 4}
CICLO_ROMANO = {"I": 1, "II": 2}

CICLO_NOMBRES: dict[str, str] = {
    "1 1": "Primer Año - I Ciclo",
    "1 2": "Primer Año - II Ciclo",
    "2 1": "Segundo Año - I Ciclo",
    "2 2": "Segundo Año - II Ciclo",
    "3 1": "Tercer Año - I Ciclo",
    "3 2": "Tercer Año - II Ciclo",
    "4 1": "Cuarto Año - I Ciclo",
    "4 2": "Cuarto Año - II Ciclo",
}

PLACEHOLDER_DASH = "—"

COMPLETAR_CICLOS_RE = re.compile(
    r"^Completar ciclos\s+(\d+\s+\d+)\s+a\s+(\d+\s+\d+)", re.IGNORECASE
)

HEADER_ESPERADO = "Nivel"


# ── Modelo ────────────────────────────────────────────────────────────────────

@dataclass
class CursoMallaExcel:
    codigo: str
    nombre: str
    creditos: float
    ciclo_id: str
    ciclo_nombre: str
    carrera_sigla: str
    carrera_nombre: str
    prerequisitos: list[str] = field(default_factory=list)
    correquisitos: list[str] = field(default_factory=list)
    es_optativo: bool = False
    prerequisito_ciclo: dict | None = None

    def to_dict(self) -> dict:
        return asdict(self)


def guardar_json(datos: list[CursoMallaExcel], ruta: Path) -> None:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump([d.to_dict() for d in datos], f, ensure_ascii=False, indent=2)
    log.info(f"  ✓ Guardado: {ruta} ({len(datos)} registros)")


# ── Helpers de parseo de celdas ─────────────────────────────────────────────────

def _normalizar_texto(valor) -> str:
    """Minúsculas, espacios colapsados y sin espacios al borde."""
    return " ".join(str(valor).split()).strip().lower()


def _es_vacio(valor) -> bool:
    if valor is None:
        return True
    texto = str(valor).strip()
    return texto == "" or texto == PLACEHOLDER_DASH


def _es_admision(valor) -> bool:
    return _normalizar_texto(valor) in ("admisión", "admision")


def _es_codigo_placeholder(codigo_raw) -> bool:
    """'—' (sin código propio) u 'OPT' (placeholder explícito de optativa)."""
    if codigo_raw is None:
        return True
    texto = str(codigo_raw).strip()
    return texto == "" or texto == PLACEHOLDER_DASH or texto.upper() == "OPT"


def _clasificar_placeholder(codigo_raw, nombre_curso: str) -> str | None:
    """
    Retorna 'GEN', 'OPT', None (no es placeholder, código real) o
    'DESCONOCIDO' (código placeholder pero el nombre no matchea ninguna
    familia conocida — caso no visto en los datos reales, se loguea y
    la fila se descarta en el caller en vez de adivinar).
    """
    if not _es_codigo_placeholder(codigo_raw):
        return None
    normalizado = _normalizar_texto(nombre_curso)
    if normalizado.startswith("estudios generales"):
        return "GEN"
    if "optativ" in normalizado:
        return "OPT"
    return "DESCONOCIDO"


def _parse_creditos(valor, contexto: str) -> float:
    if isinstance(valor, (int, float)):
        return float(valor)
    if _es_vacio(valor):
        return 0.0
    texto = str(valor).strip()
    try:
        return float(texto)
    except ValueError:
        log.warning(f"  ⚠ {contexto}: créditos no numéricos '{texto}', se usa 0")
        return 0.0


def _parse_ciclo_id(nivel, ciclo) -> tuple[str, str]:
    nivel_num = NIVEL_ROMANO[str(nivel).strip().upper()]
    ciclo_num = CICLO_ROMANO[str(ciclo).strip().upper()]
    ciclo_id = f"{nivel_num} {ciclo_num}"
    ciclo_nombre = CICLO_NOMBRES.get(ciclo_id, f"Ciclo {ciclo_id}")
    return ciclo_id, ciclo_nombre


def _split_tokens(valor) -> list[str]:
    return [t.strip() for t in str(valor).split(",") if t.strip()]


def _resolver_prerreq_tokens(tokens: list[str], nombre_a_codigo: dict[str, str]) -> list[str]:
    """
    Cada token puede ser un código de curso real, o (caso Administración
    Optativa III/IV) el nombre de una fila optativa hermana ya vista en
    la misma carrera — en ese caso se resuelve al código sintético
    generado para esa fila en vez de tratarlo como código literal.
    """
    resueltos = []
    for tok in tokens:
        clave = _normalizar_texto(tok)
        resueltos.append(nombre_a_codigo.get(clave, tok.strip().upper()))
    return resueltos


def _procesar_requisitos(valor, nombre_a_codigo: dict[str, str]) -> tuple[list[str], dict | None]:
    if _es_vacio(valor) or _es_admision(valor):
        return [], None
    texto = str(valor).strip()
    m = COMPLETAR_CICLOS_RE.match(texto)
    if m:
        return [], {"desde": m.group(1), "hasta": m.group(2)}
    return _resolver_prerreq_tokens(_split_tokens(valor), nombre_a_codigo), None


def _procesar_correquisitos(valor, contexto: str, nombre_a_codigo: dict[str, str]) -> list[str]:
    if _es_vacio(valor) or _es_admision(valor):
        return []
    texto = str(valor).strip()
    if COMPLETAR_CICLOS_RE.match(texto):
        log.warning(
            f"  ⚠ {contexto}: patrón 'Completar ciclos X a Y' inesperado en "
            f"Correquisitos ('{texto}') — se conserva como texto crudo, revisar a mano"
        )
        return [texto]
    return _resolver_prerreq_tokens(_split_tokens(valor), nombre_a_codigo)


# ── Lectura de hoja ───────────────────────────────────────────────────────────

def _leer_filas_carrera(ws) -> list[tuple]:
    """
    Ubica la fila de encabezado ('Nivel'|'Ciclo'|...) y retorna las filas
    de datos hasta la primera fila en blanco (separador antes de las notas
    al pie, que no se procesan).
    """
    filas = list(ws.iter_rows(values_only=True))
    idx_header = next(
        i for i, r in enumerate(filas)
        if r and r[0] is not None and str(r[0]).strip() == HEADER_ESPERADO
    )
    datos = []
    for r in filas[idx_header + 1:]:
        if r[0] is None:
            break
        datos.append(r)
    return datos


def _cargar_hoja(ws, sigla: str, nombre_carrera: str) -> list[CursoMallaExcel]:
    filas = _leer_filas_carrera(ws)

    gen_n = 0
    opt_n = 0
    nombre_a_codigo: dict[str, str] = {}
    crudos: list[dict] = []

    for fila in filas:
        nivel, ciclo, codigo_raw, curso, creditos_raw, req_raw, correq_raw = fila[:7]

        if curso is None or str(curso).strip() == "":
            log.warning(f"  ⚠ {sigla}: fila sin nombre de curso, se omite: {fila[:3]}")
            continue

        if str(nivel).strip() == "Opt":
            ciclo_id, ciclo_nombre = "Opt", "Optativos"
        else:
            ciclo_id, ciclo_nombre = _parse_ciclo_id(nivel, ciclo)

        familia = _clasificar_placeholder(codigo_raw, curso)
        es_optativo = False

        if familia == "GEN":
            gen_n += 1
            codigo = f"GEN{gen_n}"
            nombre_a_codigo[_normalizar_texto(curso)] = codigo
        elif familia == "OPT":
            opt_n += 1
            codigo = f"OPT{opt_n}_{sigla.upper()}"
            es_optativo = True
            nombre_a_codigo[_normalizar_texto(curso)] = codigo
        elif familia == "DESCONOCIDO":
            log.warning(
                f"  ⚠ {sigla}: código placeholder '{codigo_raw}' con nombre no "
                f"reconocido como Generales/Optativa: '{curso}' — fila omitida"
            )
            continue
        else:
            codigo = str(codigo_raw).strip().upper()

        creditos = _parse_creditos(creditos_raw, contexto=f"{sigla}/{codigo}")

        crudos.append({
            "codigo": codigo,
            "nombre": str(curso).strip(),
            "creditos": creditos,
            "ciclo_id": ciclo_id,
            "ciclo_nombre": ciclo_nombre,
            "es_optativo": es_optativo,
            "_req_raw": req_raw,
            "_correq_raw": correq_raw,
        })

    if gen_n not in (0, 4):
        log.warning(f"  ⚠ {sigla}: se generaron {gen_n} placeholders GEN, se esperaban 4")

    cursos = []
    for c in crudos:
        prerequisitos, prerequisito_ciclo = _procesar_requisitos(c["_req_raw"], nombre_a_codigo)
        correquisitos = _procesar_correquisitos(
            c["_correq_raw"], f"{sigla}/{c['codigo']}", nombre_a_codigo
        )
        cursos.append(CursoMallaExcel(
            codigo=c["codigo"],
            nombre=c["nombre"],
            creditos=c["creditos"],
            ciclo_id=c["ciclo_id"],
            ciclo_nombre=c["ciclo_nombre"],
            carrera_sigla=sigla,
            carrera_nombre=nombre_carrera,
            prerequisitos=prerequisitos,
            correquisitos=correquisitos,
            es_optativo=c["es_optativo"],
            prerequisito_ciclo=prerequisito_ciclo,
        ))
    return cursos


# ── Orquestación ──────────────────────────────────────────────────────────────

def cargar_todas_las_carreras(ruta_excel: Path) -> list[CursoMallaExcel]:
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)

    todos: list[CursoMallaExcel] = []
    for nombre_hoja, sigla in SHEET_TO_SIGLA.items():
        if nombre_hoja not in wb.sheetnames:
            log.error(f"✗ Hoja '{nombre_hoja}' no encontrada en {ruta_excel.name}")
            continue

        nombre_carrera = CARRERAS_CAMPUS_COTO.get(sigla, nombre_hoja)
        log.info(f"▶ {nombre_hoja} [{sigla}]")
        cursos = _cargar_hoja(wb[nombre_hoja], sigla, nombre_carrera)
        log.info(f"  ✓ {len(cursos)} cursos")
        todos.extend(cursos)

    return todos


def run():
    log.info("=" * 60)
    log.info("Loader de malla curricular — PLANES_DE_ESTUDIO.xlsx")
    log.info("=" * 60)
    log.info("")

    ruta = INPUT_DIR / EXCEL_FILENAME
    todos = cargar_todas_las_carreras(ruta)

    if not todos:
        log.error("✗ No se cargó ningún curso")
        return todos

    guardar_json(todos, OUTPUT_NORMALIZED_DIR / "syllabi_normalized.json")

    carreras_ok = set(c.carrera_sigla for c in todos)
    con_prereq = sum(1 for c in todos if c.prerequisitos)
    con_prereq_ciclo = sum(1 for c in todos if c.prerequisito_ciclo)
    con_correq = sum(1 for c in todos if c.correquisitos)
    optativos = sum(1 for c in todos if c.es_optativo)

    log.info(f"{'='*60}")
    log.info("✓ COMPLETADO")
    log.info(f"  Carreras:  {len(carreras_ok)}/{len(SHEET_TO_SIGLA)}")
    log.info(f"  Cursos:    {len(todos)}")
    log.info(f"  Con prerrequisitos de curso: {con_prereq}")
    log.info(f"  Con prerrequisito de rango de ciclos: {con_prereq_ciclo}")
    log.info(f"  Con correquisitos: {con_correq}")
    log.info(f"  Optativos (familia OPT* generada): {optativos}")
    log.info(f"{'='*60}")

    return todos


if __name__ == "__main__":
    run()
